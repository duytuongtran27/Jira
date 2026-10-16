import builtins
from concurrent.futures import ThreadPoolExecutor, as_completed
from rich import print
from alm_ccm.task import TaskProperty
from alm_ccm.base import reference
from urllib.parse import urlparse
from openpyxl.styles import Font, Alignment
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter
from utils.convert_lib import ConvertLib
import os
import shutil
import pandas
import pytz
import http.client
import base64
import re
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from utils.ALM_Management import ALMManagement
import time


class KPIReportHandling:

    def __init__(self, alm_url: str, userId: str, password: str) -> None:
        self.alm_url = alm_url
        self.userId = userId
        self.password = password
        self.connection = http.client.HTTPSConnection(urlparse(alm_url).netloc)
        self.headersList = {
            "User-Agent": "Thunder Client (https://www.thunderclient.com)",
            "Accept": "application/xml",
            "OSLC-Core-Version": "2.0",
        }
        self.payload = ""
        self.ALM = ALMManagement(alm_url, userId, password)
        self.entity_dict_2_query = None
        self._uri_cache = {}

    def _get_xml_response(self, http_path: str):
        """Make an authenticated GET request and return the parsed XML root."""
        credentials = f"{self.userId}:{self.password}"
        encoded_credentials = base64.b64encode(credentials.encode('utf-8')).decode('utf-8')
        headers = {
            **self.headersList,
            "Authorization": f"Basic {encoded_credentials}",
        }
        parsed = urlparse(http_path)
        conn = http.client.HTTPSConnection(parsed.netloc)
        conn.request("GET", parsed.path, self.payload, headers)
        response = conn.getresponse()
        raw = response.read()
        conn.close()
        # Strip UTF-8 BOM if present
        raw = raw.lstrip(b'\xef\xbb\xbf')
        return ET.fromstring(raw)

    def split_into_monthly_ranges(self, time_info: list) -> list[list[str]]:
        ranges = []
        if len(time_info) > 1:
            start = time_info[0]
            end = time_info[1]
            start_m, start_y = map(int, start.split("/"))
            end_m, end_y = map(int, end.split("/"))


            current_m, current_y = start_m, start_y

            while (current_y, current_m) <= (end_y, end_m):
                ranges.append([f"{current_m}/{current_y}"])
                current_m += 1
                if current_m > 12:
                    current_m = 1
                    current_y += 1
        else:
            current_m = int(time_info[0].split("/")[0])
            current_y = int(time_info[0].split("/")[1])
            ranges.append([f"{current_m}/{current_y}"])

        return ranges


    def _get_data_from_uri_cached(self, uri_path: str):
        if uri_path not in self._uri_cache:
            self._uri_cache[uri_path] = self._get_data_from_uri(uri_path)
        return self._uri_cache[uri_path]

    def _get_data_from_uri(self, uri_path: str):
        try:
            root = self._get_xml_response(uri_path)
            namespaces = {'dcterms': 'http://purl.org/dc/terms/'}
            title_element = root.find('.//dcterms:title', namespaces)
            if title_element is not None:
                return title_element.text
            return None
        except ET.ParseError as e:
            return f"XML Parsing Error: {e}"
        except http.client.HTTPException as e:
            return f"HTTP Request Error: {e}"
        except Exception as e:
            return f"An unexpected error occurred: {e}"       

    def query_kpi_wi(self, entity_id, time_info):
        selection = [
            TaskProperty.closedate,
            TaskProperty.shortid,
            TaskProperty.dcterms__type,
            TaskProperty.contributor,
            TaskProperty.title,
            TaskProperty.subject,
            TaskProperty.timespent,
            TaskProperty.plannedfor,
            TaskProperty.status,
            TaskProperty.resolution,
            TaskProperty.created,
            TaskProperty.due,
            TaskProperty.projectarea,
            TaskProperty.process__teamarea,
            TaskProperty.com_bosch_rtc_configuration_workitemtype_customattribute_tasktype
        ]
        clause_1 = TaskProperty.contributor == reference(f'{self.alm_url}/jts/users/{entity_id}')
        start_m = int(time_info[0].split("/")[0])
        start_y = int(time_info[0].split("/")[1])
        end_m = start_m + 1 if start_m != 12 else 1
        end_y = start_y if start_m != 12 else start_y + 1 
        clause_2 = TaskProperty.closedate >= datetime(start_y, start_m, 1, 0, 0, 0, 0, tzinfo=timezone.utc)
        clause_3 = TaskProperty.closedate <= datetime(end_y, end_m, 1, 0, 0, 0, 0, tzinfo=timezone.utc)
        where_clause = clause_1 & clause_2 & clause_3
        return self.ALM.query_wi(where_clause, selection, float('inf'))

    def _get_html_response(self, http_path: str) -> str:
        credentials = f"{self.userId}:{self.password}"
        encoded_credentials = base64.b64encode(credentials.encode('utf-8')).decode('utf-8')
        html_headers = {
            **self.headersList,
            "Accept": "text/html,application/xhtml+xml",
            "Authorization": f"Basic {encoded_credentials}",
        }
        parsed = urlparse(http_path)
        conn = http.client.HTTPSConnection(parsed.netloc)
        conn.request("GET", parsed.path, self.payload, html_headers)
        response = conn.getresponse()
        return response.read().decode('utf-8', errors='replace')

    def fetch_team_area(self, work_item_uri: str):
        """Fetch team area name from the work item HTML page."""
        try:
            html = self._get_html_response(work_item_uri)
            match = re.search(
                r'href="[^"]+/teamArea[^"]*">Team Area</a>\s*</dt>\s*<dd>\s*<a\s+href="[^"]+/teamareas/[^"]*">([^<]+)</a>',
                html, re.IGNORECASE
            )
            return match.group(1).strip() if match else None
        except Exception:
            return None

    def fetch_owned_by(self, work_item_uri: str):
        """Fetch Owned By name from the work item HTML page."""
        try:
            html = self._get_html_response(work_item_uri)
            match = re.search(
                r'">Owned By</a>.*?<dd>.*?<a[^>]*>([^<]+)</a>',
                html, re.IGNORECASE | re.DOTALL
            )
            return match.group(1).strip() if match else None
        except Exception:
            return None
            
    def generate_kpi_report_by_PA(self, header_df, project_area_name, time_info):       
        for entity_dict in self.entity_dict_2_query:
            try:
                query_result = []
                total_queries = 0
                user_msg = f"[INFO] KPI query for {entity_dict['Entity_Id']} - {entity_dict['Name']}"
                month_split_ranges = self.split_into_monthly_ranges(time_info)
                t_query_start = time.perf_counter()
                for month_range in month_split_ranges:
                    query = self.query_kpi_wi(entity_dict['Entity_Id'], month_range)
                    query_result.append(query)
                    total_queries = total_queries + len(query)
                t_query = time.perf_counter() - t_query_start
                print(f"\r{user_msg} - Total queries: {total_queries}")
                all_results = [r for batch in query_result for r in batch]

                def process_result(result):
                    created_dt = ConvertLib.convert_to_GMT7(result.created).replace(tzinfo=None) if hasattr(result, 'created') and result.created else None
                    resolved_dt = ConvertLib.convert_to_GMT7(result.closedate).replace(tzinfo=None) if hasattr(result, 'closedate') and result.closedate else None
                    due_dt = ConvertLib.convert_to_GMT7(result.due).replace(tzinfo=None) if hasattr(result, 'due') and result.due else None
                    planned_for_title = self._get_data_from_uri_cached(result.plannedfor.uri) if result.plannedfor else None
                    return {
                        "Project (Task)": project_area_name,
                        "Type (Task)": result.dcterms__type,
                        "Task Type (Custom) (Task)": self._get_data_from_uri_cached(result.com_bosch_rtc_configuration_workitemtype_customattribute_tasktype.uri) if result.dcterms__type == 'Task' and result.com_bosch_rtc_configuration_workitemtype_customattribute_tasktype.uri else None,
                        "Work Item ID": result.shortid,
                        "Work Item": result.title,
                        "URL (Task)": result.uri if result.uri else None,
                        "Resolution (Task)": self._get_data_from_uri_cached(result.resolution.uri) if result.resolution else None,
                        "Status (Task)": result.status if hasattr(result, 'status') else None,
                        "Owner (Task)": self.fetch_owned_by(result.uri) if result.uri else entity_dict['Name'],
                        "Creation Date (Task)": created_dt if created_dt else None,
                        "Planned For (Task)": planned_for_title,
                        "Due Date (Task)": due_dt if due_dt else None,
                        "Resolved Date (Task)": resolved_dt if resolved_dt else None,
                        "Actual Work (custom)": result.timespent,
                        "Iteration 1": planned_for_title,
                        "Start Date (Iteration 1)": created_dt if created_dt else None,
                        "End Date (Iteration 1)": resolved_dt if resolved_dt else None,
                        "Team (Task)": self.fetch_team_area(result.uri) if result.uri else None,
                        "Tags (Task)": result.subject,
                        "Actual Effort (hrs)": ConvertLib.convert_milliseconds_to_hours(result.timespent)
                    }

                t_process_start = time.perf_counter()
                with ThreadPoolExecutor(max_workers=16) as executor:
                    futures = {executor.submit(process_result, r): r for r in all_results}
                    for future in as_completed(futures):
                        header_df.append(future.result())
                t_process = time.perf_counter() - t_process_start
                print(f"[TIME] {entity_dict['Entity_Id']} | Query: {t_query:.2f}s | Process: {t_process:.2f}s | Total: {t_query + t_process:.2f}s")
            except Exception as e:
                print(f"[ERROR] Skipping {entity_dict['Entity_Id']} - {entity_dict['Name']}: {e}")
                continue

    def _apply_kpi_sheet_formatting(self, ws, df):
        header_font = Font(name='Arial', size=14, bold=True, color=Color(theme=4, type='theme'))
        header_alignment = Alignment(horizontal='left')
        ws.row_dimensions[1].height = 17.4
        for cell in ws[1]:
            cell.font = header_font
            cell.alignment = header_alignment

        url_col_idx = next((i + 1 for i, col in enumerate(df.columns) if col == "URL (Task)"), None)
        if url_col_idx:
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=url_col_idx, max_col=url_col_idx):
                cell = row[0]
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0000FF", underline="single")

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(
                (len(str(cell.value)) for cell in col if cell.value is not None),
                default=0
            )
            ws.column_dimensions[col_letter].width = max_len + 2

    def process_kpi_df(self, kpi_rows, time_info, output_path):
        if not kpi_rows:
            print("[WARNING] No KPI data to export.")
            # return

        date_cols = ["Creation Date (Task)", "Due Date (Task)", "Resolved Date (Task)", "Start Date (Iteration 1)", "End Date (Iteration 1)"]
        resolved_col = "Resolved Date (Task)"

        df = pandas.DataFrame(kpi_rows)
        for col in date_cols:
            if col in df.columns:
                df[col] = pandas.to_datetime(df[col], errors='coerce')

        kpi_filename = 'KPI_Status_Report.xlsx'

        time_info_list = self.split_into_monthly_ranges(time_info)

        with pandas.ExcelWriter(kpi_filename, engine='openpyxl') as writer:
            for month_entry in time_info_list:
                m = int(month_entry[0].split("/")[0])
                y = int(month_entry[0].split("/")[1])
                next_m = m % 12 + 1
                next_y = y + 1 if m == 12 else y

                range_start = pandas.Timestamp(y, m, 1)
                range_end = pandas.Timestamp(next_y, next_m, 1)

                if resolved_col in df.columns:
                    mask = (df[resolved_col] >= range_start) & (df[resolved_col] < range_end)
                    month_df = df[mask].copy()
                else:
                    month_df = pandas.DataFrame(columns=df.columns)

                month_label = f'{m}.{y}'
                out_df = month_df.copy()
                for col in date_cols:
                    if col in out_df.columns:
                        out_df[col] = out_df[col].dt.strftime('%m/%d/%Y')

                tl_sheet = f'TL_{month_label}'
                out_df.to_excel(writer, sheet_name=tl_sheet, index=False)
                self._apply_kpi_sheet_formatting(writer.sheets[tl_sheet], out_df)

                df_df = out_df[out_df['Type (Task)'].isin(['Problem', 'Defect'])].copy()
                df_sheet = f'DF_{month_label}'
                df_df.to_excel(writer, sheet_name=df_sheet, index=False)
                self._apply_kpi_sheet_formatting(writer.sheets[df_sheet], df_df)

        os.makedirs(output_path, exist_ok=True)
        shutil.copy(kpi_filename, output_path)
        os.remove(kpi_filename)
        print("[DONE] KPI Status report generation completed..")
